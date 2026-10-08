import csv

from openpyxl import load_workbook

from finsight.export import field_rows, write_csv, write_xlsx
from finsight.schemas import IncomeStatement, Invoice


def inc():
    return IncomeStatement(
        revenue=383285, net_income=None, unit="millions", source_pages={"revenue": 40}
    )


def test_field_rows_keep_unit_and_page():
    rows = {r[0]: r for r in field_rows(inc())}
    assert rows["revenue"] == ["revenue", 383285.0, "millions", "USD", 40]
    assert rows["net_income"][1] is None


def test_xlsx_has_sheet_per_model_and_line_sheet(tmp_path):
    invoice = Invoice(total=10, lines=[{"description": "a", "amount": 10}])
    p = tmp_path / "o.xlsx"
    write_xlsx(p, {"income": inc(), "invoice": invoice})
    wb = load_workbook(p)
    assert wb.sheetnames == ["income", "invoice", "invoice_lines"]
    assert any(r[0].value == "revenue" and r[1].value == 383285 for r in wb["income"].iter_rows())
    assert wb["invoice_lines"]["A2"].value == "a"
    vals = [c.value for row in wb["income"].iter_rows() for c in row]
    assert "null" not in vals


def test_csv(tmp_path):
    p = tmp_path / "o.csv"
    write_csv(p, Invoice(total=10, lines=[{"description": "a", "amount": 10}]))
    rows = list(csv.reader(open(p)))
    assert rows[0] == ["field", "value", "unit", "currency", "page"]
    assert ["lines"] in rows
