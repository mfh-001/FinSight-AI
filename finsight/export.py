"""CSV and XLSX export of extraction results."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from pydantic import BaseModel


def field_rows(model: BaseModel) -> list[list]:
    """One row per scalar field: name, value, unit, currency, source page."""
    d = model.model_dump()
    unit = d.get("unit", "")
    cur = d.get("currency", "")
    pages = d.get("source_pages", {}) or {}
    rows = []
    for k, v in d.items():
        if k in ("source_pages", "unit", "currency") or isinstance(v, (list, dict)):
            continue
        rows.append([k, v, unit, cur, pages.get(k)])
    return rows


def table_rows(model: BaseModel) -> dict[str, list[dict]]:
    return {k: v for k, v in model.model_dump().items() if isinstance(v, list) and v}


def write_csv(path: str | Path, model: BaseModel) -> None:
    with open(path, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["field", "value", "unit", "currency", "page"])
        w.writerows(field_rows(model))
        for name, rows in table_rows(model).items():
            w.writerow([])
            w.writerow([name])
            w.writerow(list(rows[0]))
            w.writerows([list(r.values()) for r in rows])


def write_xlsx(path: str | Path, results: dict[str, BaseModel]) -> None:
    """results maps a sheet name to a model. Lists inside a model get their own sheet."""
    wb = Workbook()
    wb.remove(wb.active)
    for name, model in results.items():
        ws = wb.create_sheet(name[:31])
        ws.append(["field", "value", "unit", "currency", "page"])
        for r in field_rows(model):
            ws.append(r)
        for sub, rows in table_rows(model).items():
            ws2 = wb.create_sheet(f"{name}_{sub}"[:31])
            ws2.append(list(rows[0]))
            for r in rows:
                ws2.append(list(r.values()))
    wb.save(path)
